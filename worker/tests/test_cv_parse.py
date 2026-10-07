"""CV parser (Phase 1.1 + OCR + AI #1 soft path)."""

from __future__ import annotations

import io
import os
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from worker.cv_parse import PARSER_VERSION, extract, extract_text, parse_bytes, parse_text
from worker.cv_parse.dates import merge_years, parse_month
from worker.cv_parse.ocr import reset_ocr_cache
from worker.cv_parse.sections import split_sections

FIXTURES = Path(__file__).parent / "fixtures" / "cv"


class CvParseTextTest(unittest.TestCase):
    def setUp(self):
        # Keep these tests on the rules/OCR path even if a local API key is set.
        self._env = patch.dict(os.environ, {"CV_AI_FALLBACK_ENABLED": "0"}, clear=False)
        self._env.start()

    def tearDown(self):
        self._env.stop()

    def test_sample_backend_extracts_contact_skills_and_years(self):
        text = (FIXTURES / "sample_backend.txt").read_text(encoding="utf-8")
        profile = parse_text(text)
        contact = profile["contact"]
        self.assertEqual(contact["full_name"], "Aysel Məmmədli")
        self.assertEqual(contact["email"], "aysel.mammadli@example.com")
        self.assertIn("+994", contact["phone"])
        self.assertIn("github.com/ayselm", contact["links"]["github"])
        self.assertIn("linkedin.com/in/aysel-mammadli", contact["links"]["linkedin"])
        self.assertEqual(contact["links"]["portfolio"], "https://aysel.dev")

        names = {item["name"] for item in profile["skills"]}
        for expected in (
            "Java",
            "Spring",
            "Kafka",
            "PostgreSQL",
            "Docker",
            "Kubernetes",
            "AWS",
        ):
            self.assertIn(expected, names)

        self.assertGreaterEqual(profile["total_years"], 5.0)
        self.assertEqual(profile["seniority"], "middle")
        self.assertIn("Java and Spring", profile.get("summary") or "")
        self.assertEqual(profile["parse_meta"]["method"], "rules")
        self.assertEqual(profile["parse_meta"]["parser_version"], PARSER_VERSION)
        self.assertGreaterEqual(profile["parse_meta"]["confidence"], 0.6)
        self.assertIn("experience", profile["parse_meta"]["sections_found"])
        self.assertIn("summary", profile["parse_meta"]["sections_found"])
        self.assertTrue(any(job.get("start") == "2021-03" for job in profile["work_history"]))
        self.assertTrue(
            any(job.get("end") is None for job in profile["work_history"]),
            profile["work_history"],
        )
        self.assertTrue(
            any("Built APIs" in (job.get("summary") or "") for job in profile["work_history"]),
            profile["work_history"],
        )
        self.assertEqual(profile["education"][0]["school"], "ADA University")
        lang_codes = {item["code"] for item in profile["languages"]}
        self.assertEqual(lang_codes, {"az", "en", "ru"})

    def test_sections_multilingual(self):
        text = """
Опыт
Engineer at X
Skills
Python
Təhsil
Baku State
"""
        sections = split_sections(text)
        self.assertIn("experience", sections)
        self.assertIn("skills", sections)
        self.assertIn("education", sections)

    def test_merge_overlapping_years(self):
        a = (date(2020, 1, 1), date(2022, 1, 1))
        b = (date(2021, 1, 1), date(2023, 1, 1))
        self.assertEqual(merge_years([a, b]), 3.0)

    def test_parse_month_present(self):
        today = date.today()
        got = parse_month("Present", end=True)
        self.assertIsNotNone(got)
        self.assertEqual(got.year, today.year)
        self.assertEqual(got.month, today.month)

    def test_parse_month_ru_abbreviation(self):
        from worker.cv_parse.dates import find_ranges

        ranges = find_ranges("Февраль 2026 – н.в.")
        self.assertEqual(len(ranges), 1)
        self.assertEqual(ranges[0][0], date(2026, 2, 1))

    def test_sample_ru_ats_jobs_and_languages(self):
        text = (FIXTURES / "sample_ru_ats.txt").read_text(encoding="utf-8")
        profile = parse_text(text)
        self.assertEqual(profile["headline"], "Java Backend Engineer")
        self.assertEqual(profile["contact"]["city"], "Баку")
        self.assertEqual(profile["contact"]["country"], "Азербайджан")
        titles = [job["title"] for job in profile["work_history"]]
        self.assertIn("Software Engineer / Backend Developer", titles)
        self.assertIn("Java Software Developer", titles)
        self.assertIn("Java Mentor", titles)
        self.assertTrue(any(job.get("company", "").startswith("Банк ВТБ") for job in profile["work_history"]))
        self.assertTrue(any(job.get("end") is None for job in profile["work_history"]))
        self.assertGreater(profile["total_years"], 0)
        lang_codes = {item["code"] for item in profile["languages"]}
        self.assertEqual(lang_codes, {"az", "en", "ru"})
        self.assertIn("experience", profile["parse_meta"]["sections_found"])
        self.assertIn("skills", profile["parse_meta"]["sections_found"])

    def test_sample_az_hazirda_dates(self):
        from worker.cv_parse.dates import find_ranges

        ranges = find_ranges("Fevral 2026 – Hazırda")
        self.assertEqual(len(ranges), 1)
        self.assertEqual(ranges[0][0], date(2026, 2, 1))

    def test_az_day_prefixed_and_ascii_i_month_ranges(self):
        from worker.cv_parse.dates import find_ranges

        afb = find_ranges('"AFB Bank" ASC | 16 Fev 2026-03 İyul 2026')
        self.assertEqual(len(afb), 1)
        self.assertEqual(afb[0][0], date(2026, 2, 1))
        self.assertEqual(afb[0][1], date(2026, 7, 31))

        # PDF often uses ASCII "I" (→ ı) instead of Turkish "İ" in "Iyul".
        intern = find_ranges("DevJoint | Iyul 2026-Avqust 2026")
        self.assertEqual(len(intern), 1)
        self.assertEqual(intern[0][0], date(2026, 7, 1))
        self.assertEqual(intern[0][1], date(2026, 8, 31))

    def test_sample_az_twocolumn_keeps_education_out_of_jobs(self):
        text = (FIXTURES / "sample_az_twocolumn.txt").read_text(encoding="utf-8")
        sections = split_sections(text)
        self.assertIn("education", sections)
        self.assertIn("skills", sections)

        profile = parse_text(text)
        titles = [job.get("title", "") for job in profile["work_history"]]
        companies = " ".join(job.get("company", "") for job in profile["work_history"])
        blob = " ".join(titles) + " " + companies
        self.assertNotIn("Universiteti", blob)
        self.assertTrue(
            any("Intern" in t or "İntern" in t for t in titles),
            titles,
        )
        self.assertTrue(
            any("AFB" in (job.get("company") or "") for job in profile["work_history"]),
            profile["work_history"],
        )
        self.assertTrue(
            any("DevJoint" in (job.get("company") or "") for job in profile["work_history"]),
            profile["work_history"],
        )
        # Intern calendar time is half-weighted; ~6 months intern ≪ 6 months senior.
        self.assertLess(profile["total_years"], 0.4)
        self.assertEqual(profile["seniority"], "intern")
        self.assertTrue(
            all(job.get("employment_type") == "internship" for job in profile["work_history"]),
            profile["work_history"],
        )
        self.assertEqual(profile["contact"]["full_name"], "SƏMA SƏFƏROVA")
        self.assertTrue(profile["education"], profile["education"])
        self.assertTrue(
            any("Universiteti" in (edu.get("school") or "") for edu in profile["education"]),
            profile["education"],
        )
        degrees = {edu.get("degree", "").lower() for edu in profile["education"]}
        self.assertTrue(degrees & {"magistr", "bakalavr"}, degrees)
        lang_codes = {item["code"] for item in profile["languages"]}
        self.assertEqual(lang_codes, {"az", "en", "tr"})

    def test_intern_years_do_not_inflate_seniority(self):
        text = """
Ada Example
ada@example.com
Experience
Java Backend Developer Intern
Trainee Soft | Jan 2018 – Dec 2023
Java Spring
Education
BSc CS
Skills
Java, Spring
"""
        profile = parse_text(text)
        self.assertEqual(profile["seniority"], "intern")
        # 6 calendar years of internship → 3 weighted years, still not "senior".
        self.assertGreaterEqual(profile["total_years"], 2.5)
        self.assertLess(profile["total_years"], 4.0)
        self.assertNotEqual(profile["seniority"], "senior")

    def test_professional_years_drive_seniority_with_intern_side(self):
        text = """
Ada Example
ada@example.com
Experience
Senior Java Developer
Acme Soft | Jan 2019 – Dec 2024
Java Spring Kafka
Java Intern
Campus Labs | Jun 2018 – Aug 2018
Java
Skills
Java, Spring, Kafka
"""
        profile = parse_text(text)
        self.assertEqual(profile["seniority"], "senior")
        self.assertGreaterEqual(profile["total_years"], 6.0)
        types = {job.get("employment_type") for job in profile["work_history"]}
        self.assertIn("internship", types)
        self.assertIn("", types)


class CvParseFilesTest(unittest.TestCase):
    def setUp(self):
        self._env = patch.dict(os.environ, {"CV_AI_FALLBACK_ENABLED": "0"}, clear=False)
        self._env.start()

    def tearDown(self):
        self._env.stop()

    def test_plain_text_bytes(self):
        data = b"Ada Lovelace\nada@example.com\nSkills\nPython, Django\n"
        profile = parse_bytes(data, filename="cv.txt")
        self.assertEqual(profile["contact"]["email"], "ada@example.com")
        names = {item["name"] for item in profile["skills"]}
        self.assertIn("Python", names)
        self.assertIn("Django", names)

    def test_docx_extract_and_parse(self):
        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("Nina Volkov")
        doc.add_paragraph("nina@example.com")
        doc.add_paragraph("Experience")
        doc.add_paragraph("Python Developer at Lab")
        doc.add_paragraph("Jan 2022 – Present")
        doc.add_paragraph("Python FastAPI PostgreSQL")
        doc.add_paragraph("Skills")
        doc.add_paragraph("Python, FastAPI, PostgreSQL")
        doc.save(buf)
        profile = parse_bytes(buf.getvalue(), filename="n.docx")
        self.assertEqual(profile["contact"]["email"], "nina@example.com")
        names = {item["name"] for item in profile["skills"]}
        self.assertIn("Python", names)
        self.assertIn("FastAPI", names)

    def test_pdf_roundtrip_extract(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=400, height=200)
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 40 120 Td (Java Spring Kafka) Tj ET")
        page[NameObject("/Contents")] = stream
        font = DictionaryObject()
        font[NameObject("/Type")] = NameObject("/Font")
        font[NameObject("/Subtype")] = NameObject("/Type1")
        font[NameObject("/BaseFont")] = NameObject("/Helvetica")
        fonts = DictionaryObject()
        fonts[NameObject("/F1")] = font
        resources = DictionaryObject()
        resources[NameObject("/Font")] = fonts
        page[NameObject("/Resources")] = resources
        out = io.BytesIO()
        writer.write(out)
        raw = out.getvalue()
        self.assertTrue(raw.startswith(b"%PDF"))
        self.assertGreaterEqual(len(PdfReader(io.BytesIO(raw)).pages), 1)
        text = extract_text(raw, filename="stack.pdf")
        self.assertIsInstance(text, str)

    def test_image_ocr_unavailable_has_error(self):
        reset_ocr_cache()
        with patch.dict(os.environ, {"CV_OCR_ENABLED": "1"}, clear=False):
            with patch("worker.cv_parse.text.ocr_enabled", return_value=False):
                with patch("worker.cv_parse.text.ocr_env_enabled", return_value=True):
                    profile = parse_bytes(b"\x89PNG\r\n\x1a\n", filename="scan.png")
        self.assertEqual(profile["parse_meta"]["confidence"], 0.0)
        self.assertEqual(profile["parse_meta"]["error"], "ocr_unavailable")

    def test_image_ocr_extracts_and_parses(self):
        sample = (FIXTURES / "sample_backend.txt").read_text(encoding="utf-8")
        pngish = b"\x89PNG\r\n\x1a\nfake"
        with patch("worker.cv_parse.text.ocr_env_enabled", return_value=True):
            with patch("worker.cv_parse.text.ocr_enabled", return_value=True):
                with patch("worker.cv_parse.text.ocr_image_bytes", return_value=sample):
                    profile = parse_bytes(pngish, filename="scan.png", content_type="image/png")
        self.assertEqual(profile["contact"]["email"], "aysel.mammadli@example.com")
        self.assertEqual(profile["parse_meta"]["text_extract"], "ocr")
        self.assertGreaterEqual(profile["parse_meta"]["confidence"], 0.6)
        names = {item["name"] for item in profile["skills"]}
        self.assertIn("Java", names)

    def test_pdf_low_text_uses_ocr_fallback(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        out = io.BytesIO()
        writer.write(out)
        raw = out.getvalue()
        ocr_body = (
            "Aysel Mammadli\naysel.mammadli@example.com\n"
            "Experience\nJava Developer at Acme\nJan 2020 – Present\n"
            "Skills\nJava, Spring, Kafka\n"
        )
        with patch("worker.cv_parse.text.ocr_env_enabled", return_value=True):
            with patch("worker.cv_parse.text.ocr_enabled", return_value=True):
                with patch("worker.cv_parse.text.ocr_pdf_bytes", return_value=ocr_body) as ocr_mock:
                    result = extract(raw, filename="scan.pdf")
                    profile = parse_bytes(raw, filename="scan.pdf")
        ocr_mock.assert_called()
        self.assertEqual(result.source, "ocr")
        self.assertIn("Java", result.text)
        self.assertEqual(profile["parse_meta"]["text_extract"], "ocr")
        self.assertIn("aysel.mammadli@example.com", profile["contact"]["email"])

    def test_digital_pdf_skips_ocr(self):
        text = "Contact\nperson@example.com\n" + ("Skills\nPython Django FastAPI PostgreSQL Docker\n" * 3)
        with patch("worker.cv_parse.text._from_pdf", return_value=text):
            with patch("worker.cv_parse.text.ocr_pdf_bytes") as ocr_mock:
                result = extract(b"%PDF-fake", filename="cv.pdf")
        ocr_mock.assert_not_called()
        self.assertEqual(result.source, "pdf")
        self.assertIn("person@example.com", result.text)


if __name__ == "__main__":
    unittest.main()
