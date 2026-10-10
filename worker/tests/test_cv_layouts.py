"""CV layout coverage: title/company/date placed differently per template style.

Text fixtures mirror what pypdf / python-docx extract from real public templates
(RenderCV Classic/Ember/Engineeringresumes/Moderncv, Deedy two-column, sb2nov ATS,
JobHire Chronological/Modern DOCX). No network, no AI.
"""

from __future__ import annotations

import io
import unittest

from worker.cv_parse import parse_bytes, parse_text

RENDERCV_CLASSIC = """John Doe
 San Francisco, CA  john.doe@email.com  rendercv.com
Welcome to RenderCV
RenderCV reads a CV written in a YAML file. This is bold, italic, and link.
Education
PhD Princeton University, Computer Science
• Thesis: Efficient Neural Architecture Search
Princeton, NJ
Sept 2018 – May 2023
Experience
Nexus AI, Co-Founder & CTO
• Built foundation model infrastructure serving 2M+ monthly API requests with 99.97%
uptime
• Raised $18M Series A led by Sequoia Capital, with participation from a16z
San Francisco, CA
June 2023 – present
2 years 10 months
NVIDIA Research, Research Intern
• Designed sparse attention mechanism reducing transformer memory footprint by 4.2x
Santa Clara, CA
May 2022 – Aug 2022
4 months
Skills
Python, Kubernetes, Docker
"""

RENDERCV_EMBER = """John Doe
john.doe@email.com
Experience
Nexus AI – San Francisco, CA June 2023 – present
Co-Founder & CTO
◆ Built foundation model infrastructure serving 2M+ monthly API requests with 99.97% uptime
NVIDIA Research – Santa Clara, CA May 2022 – Aug 2022
Research Intern
◆ Designed sparse attention mechanism reducing transformer memory footprint by 4.2x
"""

RENDERCV_ENGRES = """John Doe
john.doe@email.com
Experience
Co-Founder & CTO, Nexus AI – San Francisco, CA June 2023 – present
• Built foundation model infrastructure serving 2M+ monthly API requests with 99.97% uptime
Research Intern, NVIDIA Research – Santa Clara, CA May 2022 – Aug 2022
• Designed sparse attention mechanism reducing transformer memory footprint by 4.2x
"""

RENDERCV_MODERNCV = """John Doe
john.doe@email.com
Experience
June 2023 – present Co-Founder & CTO, Nexus AI – San Francisco, CA
• Built foundation model infrastructure serving 2M+ monthly API requests with 99.97%
uptime
May 2022 – Aug 2022 Research Intern, NVIDIA Research – Santa Clara, CA
• Designed sparse attention mechanism reducing transformer memory footprint by 4.2x
"""

DEEDY = """Debarghya Das
deedy@fb.com | 607.379.5733
EDUCATION
CORNELL UNIVERSITY
BS IN COMPUTER SCIENCE
May 2014 | Ithaca, NY
SKILLS
Java • Python • Javascript
EXPERIENCE
FACEBOOK | SOFTWARE ENGINEER
Jan 2015 - Present | New York, NY
COURSERA | KPCB FELLOW + SOFTWARE ENGINEERING INTERN
June 2014 – Sep 2014 | Mountain View, CA
• Led and shipped Yoda - the admin interface for the new Phoenix platform.
• Full-stack developer - Wrote and reviewed code for JS using Backbone, Jade,
Stylus and Require and Scala using Play
GOOGLE | SOFTWARE ENGINEERING INTERN
May 2013 – Aug 2013 | Mountain View, CA
• Worked on the YouTubeCaptions team, in Javascript and Python to plan.
"""

SB2NOV = """Sourabh Bajaj Email: sourabh@sourabhbajaj.com
sourabhbajaj.com Mobile: +1-123-456-7890
Education
• Georgia Institute of Technology Atlanta, GA
Master of Science in Computer Science; GPA: 4.00 Aug 2012 – Dec 2013
Experience
• Google Mountain View, CA
Software Engineer Oct 2016 – Present
◦ TensorFlow: TensorFlow is an open source software library for numerical computation.
• Coursera Mountain View, CA
Senior Software Engineer Jan 2014 – Oct 2016
◦ Notifications: Service for sending email, push and in-app notifications.
"""

CHRONOLOGICAL = """John T. Miller
📍 Chicago, IL
✉️ john.miller@email.com
---
Professional Summary
Detail-oriented accountant with 8+ years of experience.
---
Key Skills
Excel, QuickBooks
---
Professional Experience
Senior Accountant
Horizon Manufacturing, Chicago, IL
Jan 2025 – Present
Prepare monthly and quarterly financial statements
Accountant
Greenline Consulting, Chicago, IL
Aug 2015 – Dec 2019
Handled bookkeeping and reconciliations
---
Education
B.S. in Accounting
University of Illinois at Chicago — Graduated 2015
"""


def _jobs(text: str) -> list[tuple[str, str, str | None, str | None]]:
    prof = parse_text(text)
    return [(j["title"], j["company"], j["start"], j["end"]) for j in prof["work_history"]]


class LayoutTest(unittest.TestCase):
    def test_trailing_date_layout(self):
        jobs = _jobs(RENDERCV_CLASSIC)
        self.assertEqual(jobs[0], ("Co-Founder & CTO", "Nexus AI", "2023-06", None))
        self.assertEqual(jobs[1], ("Research Intern", "NVIDIA Research", "2022-05", "2022-08"))
        self.assertEqual(len(jobs), 2)

    def test_company_date_then_title_layout(self):
        jobs = _jobs(RENDERCV_EMBER)
        self.assertEqual(jobs[0][:2], ("Co-Founder & CTO", "Nexus AI"))
        self.assertEqual(jobs[1], ("Research Intern", "NVIDIA Research", "2022-05", "2022-08"))

    def test_title_company_location_date_same_line(self):
        jobs = _jobs(RENDERCV_ENGRES)
        self.assertEqual(jobs[0][:2], ("Co-Founder & CTO", "Nexus AI"))
        self.assertEqual(jobs[1][:2], ("Research Intern", "NVIDIA Research"))

    def test_date_first_layout(self):
        jobs = _jobs(RENDERCV_MODERNCV)
        self.assertEqual(jobs[0], ("Co-Founder & CTO", "Nexus AI", "2023-06", None))
        self.assertEqual(jobs[1][:2], ("Research Intern", "NVIDIA Research"))

    def test_company_pipe_title_all_caps(self):
        jobs = _jobs(DEEDY)
        self.assertEqual([j[0] for j in jobs], [
            "SOFTWARE ENGINEER",
            "KPCB FELLOW + SOFTWARE ENGINEERING INTERN",
            "SOFTWARE ENGINEERING INTERN",
        ])
        self.assertEqual([j[1] for j in jobs], ["FACEBOOK", "COURSERA", "GOOGLE"])

    def test_all_caps_english_headings_are_sections(self):
        prof = parse_text(DEEDY)
        self.assertIn("experience", prof["parse_meta"]["sections_found"])
        self.assertIn("education", prof["parse_meta"]["sections_found"])
        self.assertIn("skills", prof["parse_meta"]["sections_found"])

    def test_bulleted_company_line_above_title_date(self):
        jobs = _jobs(SB2NOV)
        self.assertEqual(jobs[0][0], "Software Engineer")
        self.assertTrue(jobs[0][1].startswith("Google"))
        self.assertEqual(jobs[1][0], "Senior Software Engineer")
        self.assertEqual(jobs[1][2:], ("2014-01", "2016-10"))

    def test_name_before_email_label_and_headline(self):
        prof = parse_text(SB2NOV)
        self.assertEqual(prof["contact"]["full_name"], "Sourabh Bajaj")
        self.assertEqual(prof["headline"], "Software Engineer")  # not "… Oct 2016 – Present"

    def test_location_never_comes_from_education_rows(self):
        prof = parse_text(RENDERCV_CLASSIC)
        self.assertNotIn("Princeton", prof["contact"]["city"])
        self.assertNotIn("Computer", prof["contact"]["country"])
        self.assertNotEqual(prof["contact"]["city"], "italic")

    def test_emoji_location_and_headline_from_role(self):
        prof = parse_text(CHRONOLOGICAL)
        self.assertEqual((prof["contact"]["city"], prof["contact"]["country"]), ("Chicago", "IL"))
        self.assertEqual(prof["headline"], "Senior Accountant")
        self.assertEqual(len(prof["work_history"]), 2)
        self.assertGreaterEqual(len(prof["skills"]), 0)

    def test_soft_hyphen_date_separator(self):
        text = "Experience\nDevOps Engineer Sep. 2023\u00adMar. 2024\n• Led service mesh research across clusters with Istio\n"
        prof = parse_text(text)
        self.assertEqual(prof["work_history"][0]["start"], "2023-09")
        self.assertEqual(prof["work_history"][0]["end"], "2024-03")

    def test_experience_normaliser_is_idempotent_on_standard_layout(self):
        text = "Experience\nBackend Developer\nAcme | Jan 2020 – Mar 2022\n• Built APIs with Python and Django\n"
        self.assertEqual(_jobs(text), [("Backend Developer", "Acme", "2020-01", "2022-03")])


class DocxNestedTableTest(unittest.TestCase):
    def test_docx_with_content_in_nested_table(self):
        from docx import Document

        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        cell = outer.cell(0, 0)
        cell.paragraphs[0].text = "Emma Carter"
        inner = cell.add_table(rows=1, cols=1)
        icell = inner.cell(0, 0)
        icell.paragraphs[0].text = "Professional Experience"
        icell.add_paragraph("UI/UX Designer")
        icell.add_paragraph("BrightTech Solutions — Austin, TX")
        icell.add_paragraph("Apr 2024 – Present")
        icell.add_paragraph("Redesigned core mobile app experience with Figma")
        buf = io.BytesIO()
        doc.save(buf)
        prof = parse_bytes(buf.getvalue(), filename="cv.docx", ai_fallback=False)
        self.assertEqual(prof["contact"]["full_name"], "Emma Carter")
        self.assertEqual(prof["work_history"][0]["title"], "UI/UX Designer")
        self.assertEqual(prof["work_history"][0]["start"], "2024-04")


if __name__ == "__main__":
    unittest.main()
