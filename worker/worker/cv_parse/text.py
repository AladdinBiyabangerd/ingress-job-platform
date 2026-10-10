"""Pull plain text from CV file bytes (PDF / DOCX / images / plain text).

Digital extract first; Tesseract OCR for image uploads and low-text PDFs
(plan §5.1). No AI #1.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

from worker.cv_parse.ocr import (
    MIN_DIGITAL_PDF_CHARS,
    ocr_enabled,
    ocr_unavailable_reason,
    ocr_env_enabled,
    ocr_image_bytes,
    ocr_pdf_bytes,
)

_WS = re.compile(r"[ \t]+")
_BLANK = re.compile(r"\n{3,}")

_TEXT_EXTS = {".txt", ".md", ".text"}
_PDF_EXTS = {".pdf"}
_DOCX_EXTS = {".docx"}
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".gif"}
_LEGACY_EXTS = {".doc", ".rtf", ".odt"}


@dataclass(frozen=True)
class ExtractResult:
    text: str
    source: str = ""  # text | pdf | docx | ocr | pdf+ocr
    error: str | None = None
    # OCR outcome, only set when OCR was needed (image, scan, garbled text layer):
    # used | unavailable | disabled | empty
    ocr: str | None = None


def unsupported_reason(filename: str = "", content_type: str = "") -> str | None:
    """Return a short reason when extraction is not supported, else None."""
    ext = Path(filename or "").suffix.lower()
    ctype = (content_type or "").lower()
    if ext == ".doc" or ("msword" in ctype and "openxml" not in ctype):
        return "doc_legacy_not_supported"
    if ext in {".rtf", ".odt"}:
        return "format_not_supported"
    if _is_image(ext, ctype):
        if not ocr_env_enabled():
            return "ocr_disabled"
        if not ocr_enabled():
            return "ocr_unavailable"
        return None
    if ext in _LEGACY_EXTS:
        return "format_not_supported"
    return None


def extract_text(data: bytes, filename: str = "", content_type: str = "") -> str:
    """Extract UTF-8 text from CV bytes. Empty string when nothing readable."""
    return extract(data, filename=filename, content_type=content_type).text


def extract(data: bytes, filename: str = "", content_type: str = "") -> ExtractResult:
    """Extract text and report which path produced it (digital vs OCR)."""
    if not data:
        return ExtractResult(text="", source="", error="empty_file")
    reason = unsupported_reason(filename, content_type)
    if reason:
        return ExtractResult(text="", source="", error=reason)

    ext = Path(filename or "").suffix.lower()
    ctype = (content_type or "").lower()

    if _is_image(ext, ctype):
        text = _normalize(ocr_image_bytes(data))
        if text:
            return ExtractResult(text=text, source="ocr", ocr="used")
        return ExtractResult(text="", source="ocr", error="ocr_empty", ocr="empty")

    if ext in _PDF_EXTS or "pdf" in ctype or data[:4] == b"%PDF":
        return _extract_pdf(data)

    if ext in _DOCX_EXTS or "wordprocessingml" in ctype or data[:2] == b"PK":
        text = _from_docx(data)
        if text:
            return ExtractResult(text=text, source="docx")
        # ZIP that isn't a readable docx — try nothing further.
        if ext in _DOCX_EXTS or "wordprocessingml" in ctype:
            return ExtractResult(text="", source="docx", error="docx_empty")
        # Sniffed PK without docx hint: fall through to text decode below.

    if ext in _TEXT_EXTS or ctype.startswith("text/"):
        text = _normalize(_decode(data))
        return ExtractResult(text=text, source="text", error=None if text else "text_empty")

    if data[:2] == b"PK":
        return ExtractResult(text="", source="docx", error="docx_empty")

    text = _normalize(_decode(data))
    return ExtractResult(text=text, source="text", error=None if text else "text_empty")


def _is_image(ext: str, ctype: str) -> bool:
    return ext in _IMAGE_EXTS or ctype.startswith("image/")


_READABLE = (
    (0x0000, 0x024F),  # Latin (+ extended)
    (0x0370, 0x03FF),  # Greek
    (0x0400, 0x052F),  # Cyrillic
    (0x0590, 0x06FF),  # Hebrew, Arabic
    (0x1E00, 0x1FFF),  # Latin extended additional, Greek extended
    (0x3040, 0x30FF),  # Kana
    (0x4E00, 0x9FFF),  # CJK
    (0xAC00, 0xD7AF),  # Hangul
)


def looks_garbled(text: str) -> bool:
    """True when a text layer holds many glyphs from scripts nobody writes CVs in.

    PDF fonts without a ToUnicode map decode to Georgian / Tibetan / Sinhala / PUA noise.
    Real Latin, Cyrillic, CJK, Arabic, Hebrew or Greek text never trips this.
    """
    letters = [c for c in text if c.isalpha()]
    if len(letters) < 60:
        return False
    bad = sum(1 for c in letters if not any(lo <= ord(c) <= hi for lo, hi in _READABLE))
    return bad / len(letters) > 0.15


def _extract_pdf(data: bytes) -> ExtractResult:
    """Digital text layer first; OCR (optional) only for scans / garbled / empty layers."""
    digital = _from_pdf(data)
    garbled = digital if looks_garbled(digital) else ""
    if garbled:
        digital = ""  # let OCR try; keep the noise only when OCR cannot replace it
    if len(digital) >= MIN_DIGITAL_PDF_CHARS:
        return ExtractResult(text=digital, source="pdf")

    # Layer is empty, thin or garbled: OCR is the optional rescue. Failure never breaks parsing.
    if not ocr_env_enabled():
        status = "disabled"
    elif not ocr_enabled():
        status = "unavailable"
    else:
        status = ""
    if not status:
        try:
            ocr_text = _normalize(ocr_pdf_bytes(data))
        except Exception:  # OCR is optional
            ocr_text = ""
        if ocr_text and not looks_garbled(ocr_text):
            return ExtractResult(text=ocr_text, source="pdf+ocr" if digital else "ocr", ocr="used")
        status = "empty"
    if garbled:
        return ExtractResult(text=garbled, source="pdf", error="text_garbled", ocr=status)
    if digital:
        return ExtractResult(text=digital, source="pdf", ocr=status)
    return ExtractResult(text="", source="pdf" if status != "empty" else "ocr", error=f"ocr_{status}", ocr=status)



def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


_GLYPH_DIGITS = {
    "zero": "0", "one": "1", "two": "2", "three": "3", "four": "4",
    "five": "5", "six": "6", "seven": "7", "eight": "8", "nine": "9",
}
# LaTeX/Typst PDFs without ToUnicode leak glyph names: "/two.lnum", "/A.sc", "/f_ic".
_GLYPH_DIGIT = re.compile(r"/(zero|one|two|three|four|five|six|seven|eight|nine)\.(?:lnum|onum|tnum|pnum)")
_GLYPH_SC = re.compile(r"/([A-Za-z])\.sc")
_GLYPH_LIG = re.compile(r"/f_(f_)?([ilf])(?![a-z]\.)")
_GLYPH_ICON = re.compile(r"/(?:_\d{2,4}|[a-z]{2,12}_[a-z_]{2,12}|f1ab|github|twitter|linkedin)(?=[A-Za-z+@])")
_PUA = re.compile("[\ue000-\uf8ff\ufffd]")
_KERN_GAP = re.compile(r"\b([TVWYK]) (?=[a-z]{2,})")
_HYPHEN_WRAP = re.compile(r"(?<=[a-z]{2})-\n(?=[a-z]{2})")
_LIGS = {"\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi", "\ufb04": "ffl"}


def _repair_glyphs(text: str) -> str:
    """Undo common PDF-extraction artefacts (glyph names, ligatures, kerning gaps, icon fonts)."""
    if "/" in text:
        text = _GLYPH_DIGIT.sub(lambda m: _GLYPH_DIGITS[m.group(1)], text)
        text = _GLYPH_SC.sub(lambda m: m.group(1), text)
        text = _GLYPH_LIG.sub(lambda m: "f" + ("f" if m.group(1) else "") + m.group(2), text)
        text = _GLYPH_ICON.sub("", text)
    for lig, rep in _LIGS.items():
        text = text.replace(lig, rep)
    text = _PUA.sub("", text)
    text = _HYPHEN_WRAP.sub("", text)
    text = re.sub(r"\b([A-Z][a-z]{2}) \.(?=\s*(?:19|20)\d{2})", r"\1.", text)  # "Mar . 2020"
    text = re.sub(r"((?:19|20)\d{2})(?=[A-Z][a-z])", r"\1 ", text)  # "2014 - 2016IT Consultant"
    return _KERN_GAP.sub(r"\1", text)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ").replace("\ufeff", "").replace("\xad", "-")
    text = _repair_glyphs(text)
    lines = [_WS.sub(" ", line).rstrip() for line in text.split("\n")]
    return _BLANK.sub("\n\n", "\n".join(lines)).strip()


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            chunk = page.extract_text() or ""
        except Exception:
            chunk = ""
        if chunk.strip():
            parts.append(chunk)
    return _normalize("\n".join(parts))


def _from_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        return ""
    parts: list[str] = []
    for para in doc.paragraphs:
        line = (para.text or "").strip()
        if line:
            parts.append(line)
    for table in doc.tables:
        for row in table.rows:
            cells = [((cell.text or "").strip()) for cell in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    base = _normalize("\n".join(parts))
    # Nested tables / text boxes (sidebar templates) are invisible to doc.paragraphs;
    # when a document-order walk finds clearly more text, prefer it.
    full = _normalize("\n".join(_docx_all_paragraphs(doc)))
    if len(full) > len(base) * 1.15:
        return full
    return base


def _docx_all_paragraphs(doc) -> list[str]:
    """Every paragraph in document order, including nested tables and text boxes.

    mc:Fallback copies of text boxes are skipped so each box is read once.
    """
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    out: list[str] = []
    try:
        for para in doc.element.body.iter(f"{ns}p"):
            anc = para.getparent()
            skip = False
            while anc is not None:
                if anc.tag == f"{mc}Fallback":
                    skip = True
                    break
                anc = anc.getparent()
            if skip:
                continue
            chunks: list[str] = []
            for node in para.iter():
                if node.tag == f"{ns}t":
                    chunks.append(node.text or "")
                elif node.tag in (f"{ns}tab",):
                    chunks.append(" ")
                elif node.tag == f"{ns}br":
                    chunks.append("\n")
            text = "".join(chunks)
            for line in text.split("\n"):
                if line.strip():
                    out.append(line.strip())
    except Exception:
        return []
    return out
